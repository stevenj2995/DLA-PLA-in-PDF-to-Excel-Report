
from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from . import config

# membaca struktur file standar

@dataclass
class Kolom:
    indeks: int # 1-based
    huruf: str # "A", "AB", ...
    grup: str | None # header baris 2
    header: str | None # header baris 3 -> acuan pencocokan
    flag: str | None # "Mandatory" / "Optional"
    peran: str = "kosong"
    format_angka: str = "@"

    @property
    def nama(self) -> str:
        return (self.header or self.grup or self.huruf).strip()

    @property
    def nama_bersih(self) -> str:
        n = self.nama.split("\n")[0]
        return " ".join(n.replace("(YYYY-MM-DD)", "").replace("(HH:MM)", "")
                         .replace("(Y/N)", "").split()).strip()

    @property
    def nama_lengkap(self) -> str:
        if self.grup and self.header and self.grup != self.header:
            return f"{self.grup} - {self.nama_bersih}"
        return self.nama_bersih


@dataclass
class Skema:
    kolom: list[Kolom] = field(default_factory=list)

    def __iter__(self):
        return iter(self.kolom)

    def __len__(self):
        return len(self.kolom)

    @property
    def target_pencocokan(self) -> list[Kolom]:
        # kolom yang boleh jadi tujuan pencocokan parameter PDF
        return [k for k in self.kolom if k.peran == "dari_pdf"]

    def ringkasan(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for k in self.kolom:
            out[k.peran] = out.get(k.peran, 0) + 1
        return out


def _peran(huruf: str) -> str:
    if huruf == config.KOLOM_NOMOR_URUT:
        return "nomor"
    if huruf in config.KOLOM_KONSTAN:
        return "konstan"
    if huruf in config.KOLOM_RUMUS:
        return "rumus"
    if huruf in config.KOLOM_DARI_PDF:
        return "dari_pdf"
    if huruf in config.KOLOM_MONITORING:
        return "monitoring"
    if huruf in config.KOLOM_FEE:
        return "fee"
    return "kosong"


@lru_cache(maxsize=4)
def muat(path_standar: str | None = None) -> Skema:
    path = Path(path_standar) if path_standar else config.file_standar()
    wb = openpyxl.load_workbook(path, read_only=False)
    ws = wb[config.SHEET_UTAMA]

    grup: dict[int, str] = {}
    for m in ws.merged_cells.ranges:
        if m.min_row == config.BARIS_GRUP:
            nilai = ws.cell(config.BARIS_GRUP, m.min_col).value
            for c in range(m.min_col, m.max_col + 1):
                if nilai is not None:
                    grup[c] = str(nilai).strip()

    kolom: list[Kolom] = []
    for c in range(1, ws.max_column + 1):
        huruf = get_column_letter(c)
        g = grup.get(c)
        if g is None:
            v = ws.cell(config.BARIS_GRUP, c).value
            g = str(v).strip() if v is not None else None
        h = ws.cell(config.BARIS_HEADER, c).value
        f = ws.cell(config.BARIS_FLAG, c).value
        fmt = ws.cell(config.BARIS_DATA_AWAL, c).number_format or "@"
        kolom.append(Kolom(
            indeks=c, huruf=huruf, grup=g,
            header=str(h).strip() if h is not None else None,
            flag=str(f).strip() if f is not None else None,
            peran=_peran(huruf), format_angka=fmt,
        ))
    wb.close()
    return Skema(kolom)

# menulis hasil ke file Excel

RE_EXTLST = re.compile(r"<extLst>.*?</extLst>", re.S)
RE_XR_UID = re.compile(r'\s+xr:uid="[^"]*"')
RE_X14_DV = re.compile(r"<x14:dataValidation\b")


@dataclass
class Baris:
    nilai: dict[str, object] = field(default_factory=dict) # kunci = huruf kolom
    sumber: str = "" # nama file PDF asalnya
    peringatan: list[str] = field(default_factory=list)


# nama part XML untuk sheet MosyClaimTask
def _sheet_xml_utama(zf: zipfile.ZipFile) -> str:
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    m = re.search(
        r'<sheet[^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(config.SHEET_UTAMA),
        wb_xml,
    )
    if not m:
        return "xl/worksheets/sheet1.xml"
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    m2 = re.search(r'<Relationship[^>]*Id="%s"[^>]*Target="([^"]+)"' % m.group(1), rels)
    if not m2:
        return "xl/worksheets/sheet1.xml"
    target = m2.group(1).lstrip("/")
    return target if target.startswith("xl/") else f"xl/{target}"


def _ambil_extlst(path_standar: Path, part: str) -> str | None:
    with zipfile.ZipFile(path_standar) as z:
        xml = z.read(part).decode("utf-8")
    m = RE_EXTLST.search(xml)
    if not m:
        return None
    return RE_XR_UID.sub("", m.group(0))


def _tambal_extlst(sumber: Path, tujuan: Path, part: str, blok: str) -> None:
    with zipfile.ZipFile(sumber) as zin, \
         zipfile.ZipFile(tujuan, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                s = RE_EXTLST.sub("", data.decode("utf-8"))
                s = s.replace("</worksheet>", blok + "</worksheet>")
                data = s.encode("utf-8")
            zout.writestr(item, data)

def hitung_dropdown(path: Path) -> int:
    with zipfile.ZipFile(path) as z:
        xml = z.read(_sheet_xml_utama(z)).decode("utf-8", "ignore")
    return len(RE_X14_DV.findall(xml))


def tulis(
    baris: list[Baris],
    tujuan: Path,
    *,
    email_operator: str,
    path_standar: Path | None = None,
    skema: Skema | None = None,
    mulai_nomor: int = 1,
) -> dict:
    path_standar = Path(path_standar or config.file_standar())
    skema = skema or muat(str(path_standar))
    tujuan = Path(tujuan)
    tujuan.parent.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(path_standar) as z:
        part = _sheet_xml_utama(z)
    blok = _ambil_extlst(path_standar, part)
    dropdown_asli = hitung_dropdown(path_standar)

    wb = openpyxl.load_workbook(path_standar)
    ws = wb[config.SHEET_UTAMA]

    fmt = {k.huruf: ws.cell(config.BARIS_DATA_AWAL, k.indeks).number_format for k in skema}

    for r in range(config.BARIS_DATA_AWAL, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None

    konstan = dict(config.KOLOM_KONSTAN)
    konstan[config.KOLOM_EMAIL_OPERATOR] = email_operator

    for i, b in enumerate(baris):
        r = config.BARIS_DATA_AWAL + i
        for k in skema:
            sel = ws.cell(r, k.indeks)
            sel.number_format = fmt.get(k.huruf, "@")

            if k.huruf == config.KOLOM_NOMOR_URUT:
                sel.value = str(mulai_nomor + i)
            elif k.huruf in konstan and k.huruf not in b.nilai:
                v = konstan[k.huruf]
                if v is not None:
                    sel.value = v
            elif k.huruf in config.KOLOM_RUMUS:
                sel.value = _rumus(k.huruf, r, b)
            elif k.huruf in b.nilai:
                v = b.nilai[k.huruf]
                sel.value = None if v is None or v == "" else v

    wb.save(tujuan)
    wb.close()

    dropdown_setelah = hitung_dropdown(tujuan)
    if blok and dropdown_setelah < dropdown_asli:
        sementara = tujuan.with_suffix(".tmp.xlsx")
        shutil.move(str(tujuan), str(sementara))
        try:
            _tambal_extlst(sementara, tujuan, part, blok)
        finally:
            sementara.unlink(missing_ok=True)
        dropdown_setelah = hitung_dropdown(tujuan)

    return {
        "file": str(tujuan),
        "baris": len(baris),
        "dropdown_asli": dropdown_asli,
        "dropdown_hasil": dropdown_setelah,
        "dropdown_utuh": dropdown_setelah == dropdown_asli,
    }


def _rumus(huruf: str, r: int, b: Baris) -> str | None:
    pola = config.KOLOM_RUMUS[huruf]
    if huruf == "AR":
        share = b.nilai.get("_share_aab")
        return pola.format(r=r, share=share) if share else None
    return pola.format(r=r)
